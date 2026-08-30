"""
ETL: Extrae datos de ingresos desde PDC_Ppto_Comercial_2026 y Consolidado_Ticketing
Fuentes:
  - bc117802-PDC_Ppto_Comercial_2026_Ingresos.xlsx
  - ee07e43d-Consolidado_Ticketing_20222026.xlsx
Salida: data/processed/ingresos.json
"""

import json
from pathlib import Path
import openpyxl

UPLOAD_DIR = Path("/root/.claude/uploads/e6ab2290-85dc-5766-b882-e4947c74218b")
SOURCE_PDC = UPLOAD_DIR / "bc117802-PDC_Ppto_Comercial_2026_Ingresos.xlsx"
SOURCE_CONSOL = UPLOAD_DIR / "ee07e43d-Consolidado_Ticketing_20222026.xlsx"
OUT = Path(__file__).parent.parent / "processed" / "ingresos.json"

MESES = ["Enero","Febrero","Marzo","Abril","Mayo","Junio",
         "Julio","Agosto","Septiembre","Octubre","Noviembre","Diciembre"]

def to_num(val):
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val) if val == val else 0.0
    s = str(val).strip().replace(".","").replace(",",".")
    try:
        return float(s)
    except ValueError:
        return 0.0

def find_header_row(ws, keyword, max_row=40):
    for row in ws.iter_rows(min_row=1, max_row=max_row):
        for cell in row:
            if isinstance(cell.value, str) and keyword.lower() in cell.value.lower():
                return cell.row
    return None

def extract_pdc_ic(ws):
    lines = []
    for row in ws.iter_rows(min_row=1, max_row=300, values_only=True):
        if not row or not row[0]:
            continue
        label = str(row[0]).strip()
        if not label or label.lower() in ("descripcion","descripción","linea","línea"):
            continue
        nums = []
        for v in row[1:]:
            if isinstance(v, (int, float)) and v == v:
                nums.append(float(v))
        if nums:
            lines.append({
                "linea": label,
                "valores": nums
            })
    return lines

def extract_locales(ws):
    locatarios = []
    header_row = find_header_row(ws, "locatario") or find_header_row(ws, "local") or 3
    for row in ws.iter_rows(min_row=header_row+1, max_row=300, values_only=True):
        if not row or not row[0]:
            continue
        label = str(row[0]).strip()
        if not label or "total" in label.lower():
            continue
        meses_vals = []
        for i, m in enumerate(MESES):
            col_idx = i + 1
            val = to_num(row[col_idx]) if col_idx < len(row) else 0.0
            meses_vals.append({"mes": m, "valor": val})
        total = sum(v["valor"] for v in meses_vals)
        if total > 0 or label:
            locatarios.append({
                "locatario": label,
                "meses": meses_vals,
                "total": total
            })
    return locatarios

def extract_espacios(ws):
    espacios = []
    header_row = find_header_row(ws, "espacio") or find_header_row(ws, "tipo") or 3
    for row in ws.iter_rows(min_row=header_row+1, max_row=300, values_only=True):
        if not row or not row[0]:
            continue
        label = str(row[0]).strip()
        if not label or "total" in label.lower():
            continue
        tipo = str(row[1]).strip() if len(row) > 1 and row[1] else "Comercial"
        meses_vals = []
        for i, m in enumerate(MESES):
            col_idx = i + 2
            val = to_num(row[col_idx]) if col_idx < len(row) else 0.0
            meses_vals.append({"mes": m, "valor": val})
        total = sum(v["valor"] for v in meses_vals)
        espacios.append({
            "espacio": label,
            "tipo": tipo,
            "meses": meses_vals,
            "total": total
        })
    return espacios

def extract_auspicios(ws):
    auspicios = []
    header_row = find_header_row(ws, "auspicio") or find_header_row(ws, "empresa") or 3
    for row in ws.iter_rows(min_row=header_row+1, max_row=300, values_only=True):
        if not row or not row[0]:
            continue
        label = str(row[0]).strip()
        if not label or "total" in label.lower():
            continue
        meses_vals = []
        for i, m in enumerate(MESES):
            col_idx = i + 1
            val = to_num(row[col_idx]) if col_idx < len(row) else 0.0
            meses_vals.append({"mes": m, "valor": val})
        total = sum(v["valor"] for v in meses_vals)
        auspicios.append({
            "auspiciador": label,
            "meses": meses_vals,
            "total": total
        })
    return auspicios

def extract_consolidado(wb_consol):
    result = {"historico": [], "operacional_2026": []}

    for sn in wb_consol.sheetnames:
        if "operacional" in sn.lower() or "ingreso" in sn.lower():
            ws = wb_consol[sn]
            for row in ws.iter_rows(min_row=1, max_row=200, values_only=True):
                if not row or not row[0]:
                    continue
                label = str(row[0]).strip()
                if not label:
                    continue
                nums = [to_num(v) for v in row[1:] if v is not None]
                if nums:
                    result["operacional_2026"].append({
                        "linea": label,
                        "valores": nums[:3]
                    })
            break

    for sn in wb_consol.sheetnames:
        if "ticket" in sn.lower():
            ws = wb_consol[sn]
            for row in ws.iter_rows(min_row=1, max_row=200, values_only=True):
                if not row or not row[0]:
                    continue
                label = str(row[0]).strip()
                nums = [to_num(v) for v in row[1:] if v is not None and isinstance(v,(int,float))]
                if nums:
                    result["historico"].append({
                        "concepto": label,
                        "series": nums
                    })
            break

    return result

def main():
    wb_pdc = openpyxl.load_workbook(SOURCE_PDC, data_only=True)
    wb_consol = openpyxl.load_workbook(SOURCE_CONSOL, data_only=True)

    print(f"PDC hojas: {wb_pdc.sheetnames}")
    print(f"Consolidado hojas: {wb_consol.sheetnames}")

    output = {
        "meta": {
            "fuentes": [SOURCE_PDC.name, SOURCE_CONSOL.name],
            "unidad": "M$ (miles de pesos chilenos)",
            "anio": 2026
        },
        "pdc_resumen": [],
        "locatarios": [],
        "espacios": [],
        "auspicios": [],
        "consolidado": {}
    }

    for sn in wb_pdc.sheetnames:
        if "pdc" in sn.lower() and "ic" in sn.lower():
            print(f"  PDC IC: {sn}")
            output["pdc_resumen"] = extract_pdc_ic(wb_pdc[sn])
            break

    for sn in wb_pdc.sheetnames:
        if "local" in sn.lower():
            print(f"  Locales: {sn}")
            output["locatarios"] = extract_locales(wb_pdc[sn])
            break

    for sn in wb_pdc.sheetnames:
        if "espacio" in sn.lower():
            print(f"  Espacios: {sn}")
            output["espacios"] = extract_espacios(wb_pdc[sn])
            break

    for sn in wb_pdc.sheetnames:
        if "auspi" in sn.lower():
            print(f"  Auspicios: {sn}")
            output["auspicios"] = extract_auspicios(wb_pdc[sn])
            break

    print(f"  Consolidado...")
    output["consolidado"] = extract_consolidado(wb_consol)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(output, f, ensure_ascii=False, indent=2)

    print(f"\nOK -> {OUT}")
    print(f"  Locatarios: {len(output['locatarios'])}")
    print(f"  Espacios: {len(output['espacios'])}")
    print(f"  Auspicios: {len(output['auspicios'])}")
    print(f"  Consolidado líneas: {len(output['consolidado'].get('operacional_2026',[]))}")

if __name__ == "__main__":
    main()
