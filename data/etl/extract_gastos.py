"""
ETL: Extrae datos de gastos desde Ppto_GAM_2026 y genera gastos.json
Fuente: dc4312ef-Ppto_GAM_2026___rea_y_CC__V19052026.xlsx
Salida: data/processed/gastos.json
"""

import json
import re
from pathlib import Path
import openpyxl

UPLOAD_DIR = Path("/root/.claude/uploads/e6ab2290-85dc-5766-b882-e4947c74218b")
SOURCE = UPLOAD_DIR / "dc4312ef-Ppto_GAM_2026___rea_y_CC__V19052026.xlsx"
OUT = Path(__file__).parent.parent / "processed" / "gastos.json"

MESES = ["Enero","Febrero","Marzo","Abril","Mayo","Junio",
         "Julio","Agosto","Septiembre","Octubre","Noviembre","Diciembre"]

AREA_SHEETS = {
    "Adm y Finanzas":       "Adm y Finanzas",
    "Programación":         "Programación",
    "Técnica":              "Técnica",
    "Producción":           "Producción",
    "Comercial y Mkt":      "Comercial y Mkt",
    "Comunicaciones":       "Comunicaciones",
    "Dirección Ejecutiva":  "Dirección Ejecutiva",
    "GxD":                  "GxD",
    "Inversiones":          "Inversiones",
}

def to_num(val):
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val) if val == val else 0.0
    s = str(val).strip().replace(".","").replace(",",".")
    try:
        return float(s)
    except ValueError:
        return 0.0

def parse_area_sheet(ws):
    header_row = None
    for row in ws.iter_rows(min_row=1, max_row=30):
        for cell in row:
            if isinstance(cell.value, str) and "enero" in cell.value.lower():
                header_row = cell.row
                break
        if header_row:
            break

    if not header_row:
        return []

    month_cols = {}
    row_cells = list(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=False))[0]
    current_mes_idx = -1
    for cell in row_cells:
        v = cell.value
        if isinstance(v, str):
            for i, m in enumerate(MESES):
                if m.lower() in v.lower():
                    current_mes_idx = i
                    if i not in month_cols:
                        month_cols[i] = {"real_col": cell.column, "ppto_col": None}
                    else:
                        month_cols[i]["ppto_col"] = cell.column
                    break

    month_col_pairs = []
    for cell in row_cells:
        v = cell.value
        if isinstance(v, str):
            for m in MESES:
                if m.lower() in v.lower():
                    month_col_pairs.append({"mes": m, "col_start": cell.column})
                    break

    items = []
    current_item = None
    current_subitems = []

    for row in ws.iter_rows(min_row=header_row + 1, values_only=False):
        label_cell = row[0] if row else None
        if not label_cell or label_cell.value is None:
            continue
        label = str(label_cell.value).strip()
        if not label or label.lower() in ("real","presupuesto","ppto","total","subtotal",""):
            continue

        is_bold = label_cell.font and label_cell.font.bold
        indent = label_cell.alignment.indent if label_cell.alignment else 0

        meses_data = []
        for pair in month_col_pairs:
            col = pair["col_start"] - 1
            real_val = to_num(row[col].value) if col < len(row) else 0.0
            ppto_val = to_num(row[col+1].value) if col+1 < len(row) else 0.0
            meses_data.append({
                "mes": pair["mes"],
                "real": real_val,
                "ppto": ppto_val
            })

        total_real = sum(m["real"] for m in meses_data)
        total_ppto = sum(m["ppto"] for m in meses_data)

        entry = {
            "nombre": label,
            "meses": meses_data,
            "total_real": total_real,
            "total_ppto": total_ppto,
            "pct_ejecucion": round(total_real / total_ppto, 4) if total_ppto else 0.0
        }

        if is_bold or indent == 0:
            if current_item:
                current_item["subitems"] = current_subitems
                items.append(current_item)
            current_item = entry
            current_item["subitems"] = []
            current_subitems = []
        else:
            current_subitems.append(entry)

    if current_item:
        current_item["subitems"] = current_subitems
        items.append(current_item)

    return items


def main():
    wb = openpyxl.load_workbook(SOURCE, data_only=True)
    print(f"Hojas disponibles: {wb.sheetnames}")

    output = {
        "meta": {
            "fuente": SOURCE.name,
            "unidad": "M$ (miles de pesos chilenos)",
            "anio": 2026
        },
        "areas": []
    }

    for area_name, sheet_name in AREA_SHEETS.items():
        ws = None
        for sn in wb.sheetnames:
            if sheet_name.lower().replace(" ","") in sn.lower().replace(" ",""):
                ws = wb[sn]
                break
        if not ws:
            print(f"  [WARN] Hoja no encontrada: {sheet_name}")
            continue

        print(f"  Procesando: {area_name} ({ws.title})")
        items = parse_area_sheet(ws)

        total_real = sum(i["total_real"] for i in items)
        total_ppto = sum(i["total_ppto"] for i in items)

        output["areas"].append({
            "area": area_name,
            "total_real": total_real,
            "total_ppto": total_ppto,
            "pct_ejecucion": round(total_real / total_ppto, 4) if total_ppto else 0.0,
            "items": items
        })

    OUT.parent.mkdir(parents=True, exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(output, f, ensure_ascii=False, indent=2)
    print(f"\nOK -> {OUT}")
    print(f"Áreas extraídas: {len(output['areas'])}")
    total_items = sum(len(a["items"]) for a in output["areas"])
    print(f"Ítems totales: {total_items}")


if __name__ == "__main__":
    main()
