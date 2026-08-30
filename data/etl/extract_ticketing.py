"""
ETL: Extrae datos de ticketing desde Liquidacion_Mensual y Reporte_Ticketing_2026
Fuentes:
  - bd9f74b4-Liquidacion_Mensual_GAM_Mensual_7_2026.xlsx
  - e68c1c20-Reporte_Ticketing_2026_2.xlsx
Salida: data/processed/ticketing.json
"""

import json
from pathlib import Path
from collections import defaultdict
import openpyxl

UPLOAD_DIR = Path("/root/.claude/uploads/e6ab2290-85dc-5766-b882-e4947c74218b")
SOURCE_LIQ = UPLOAD_DIR / "bd9f74b4-Liquidacion_Mensual_GAM_Mensual_7_2026.xlsx"
SOURCE_REP = UPLOAD_DIR / "e68c1c20-Reporte_Ticketing_2026_2.xlsx"
OUT = Path(__file__).parent.parent / "processed" / "ticketing.json"

def to_num(val):
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val) if val == val else 0.0
    s = str(val).strip().replace(".","").replace(",",".")
    try:
        return float(s)
    except ValueError:
        return 0.0

def find_header_row_with_col(ws, keywords, max_row=10):
    for row in ws.iter_rows(min_row=1, max_row=max_row):
        vals = [str(c.value).strip().lower() if c.value else "" for c in row]
        if any(kw.lower() in " ".join(vals) for kw in keywords):
            return row[0].row, {v: i for i, v in enumerate(vals) if v}
    return None, {}

def extract_tickets(ws):
    header_row, col_map = find_header_row_with_col(ws, ["event","evento","ticket"])
    if header_row is None:
        header_row = 1
        row_vals = [str(c.value).strip().lower() if c.value else "" for c in list(ws.iter_rows(min_row=1, max_row=1))[0]]
        col_map = {v: i for i, v in enumerate(row_vals) if v}

    def col(name, row_data):
        for k, idx in col_map.items():
            if name.lower() in k:
                return row_data[idx] if idx < len(row_data) else None
        return None

    by_event = defaultdict(lambda: {
        "tickets": 0, "cortesias": 0,
        "ingresos": 0.0, "comisiones": 0.0,
        "canales": defaultdict(int),
        "tipos_ticket": defaultdict(float)
    })

    rows_read = 0
    for row in ws.iter_rows(min_row=header_row+1, max_row=6000, values_only=True):
        if not row or not row[0]:
            continue
        rows_read += 1

        event = str(col("event", row) or "").strip()
        ticket_name = str(col("ticket_name", row) or col("ticket", row) or "").strip()
        channel = str(col("channel_name", row) or col("channel", row) or "").strip()
        amount = to_num(col("ticket_amount", row) or col("amount", row))
        charge = to_num(col("ticket_charge", row) or col("charge", row))
        commission = to_num(col("commission", row))
        state = str(col("state", row) or "").strip()

        if not event:
            continue

        if state and state.lower() in ("cancelled","cancelado","anulado"):
            continue

        is_cortesia = (amount == 0 and charge == 0)

        by_event[event]["tickets"] += 1
        if is_cortesia:
            by_event[event]["cortesias"] += 1
        by_event[event]["ingresos"] += charge
        by_event[event]["comisiones"] += commission
        if channel:
            by_event[event]["canales"][channel] += 1
        if ticket_name:
            by_event[event]["tipos_ticket"][ticket_name] += charge

    eventos = []
    for event_name, data in by_event.items():
        total_tickets = data["tickets"]
        ingresos = data["ingresos"]
        comisiones = data["comisiones"]
        pct_cortesias = round(data["cortesias"] / total_tickets, 4) if total_tickets else 0.0

        canales = [{"canal": k, "cantidad": v} for k, v in data["canales"].items()]
        canales.sort(key=lambda x: -x["cantidad"])

        tipos = [{"tipo": k, "ingreso": round(v, 2)} for k, v in data["tipos_ticket"].items()]
        tipos.sort(key=lambda x: -x["ingreso"])

        eventos.append({
            "evento": event_name,
            "total_tickets": total_tickets,
            "cortesias": data["cortesias"],
            "pct_cortesias": pct_cortesias,
            "ingresos_netos": round(ingresos, 2),
            "comisiones": round(comisiones, 2),
            "canales": canales,
            "tipos_ticket": tipos
        })

    eventos.sort(key=lambda x: -x["ingresos_netos"])
    print(f"    Filas leídas: {rows_read} | Eventos únicos: {len(eventos)}")
    return eventos

def extract_pagos(ws):
    header_row, col_map = find_header_row_with_col(ws, ["payment","pago","gateway"])
    if header_row is None:
        header_row = 1
        row_vals = [str(c.value).strip().lower() if c.value else "" for c in list(ws.iter_rows(min_row=1, max_row=1))[0]]
        col_map = {v: i for i, v in enumerate(row_vals) if v}

    def col(name, row_data):
        for k, idx in col_map.items():
            if name.lower() in k:
                return row_data[idx] if idx < len(row_data) else None
        return None

    by_gateway = defaultdict(lambda: {"count": 0, "total": 0.0})
    by_channel = defaultdict(lambda: {"count": 0, "total": 0.0})
    total_amount = 0.0
    rows_read = 0

    for row in ws.iter_rows(min_row=header_row+1, max_row=5000, values_only=True):
        if not row or not row[0]:
            continue
        rows_read += 1

        gateway = str(col("gateway", row) or col("payment_gateway", row) or "Desconocido").strip()
        amount = to_num(col("amount", row))
        boxoffice = col("boxoffice", row)
        channel = "Boxoffice" if boxoffice else "Internet"

        by_gateway[gateway]["count"] += 1
        by_gateway[gateway]["total"] += amount
        by_channel[channel]["count"] += 1
        by_channel[channel]["total"] += amount
        total_amount += amount

    gateways = [
        {"gateway": k, "transacciones": v["count"], "total": round(v["total"], 2)}
        for k, v in by_gateway.items()
    ]
    gateways.sort(key=lambda x: -x["total"])

    channels = [
        {"canal": k, "transacciones": v["count"], "total": round(v["total"], 2)}
        for k, v in by_channel.items()
    ]

    print(f"    Pagos leídos: {rows_read} | Total M$: {round(total_amount,2)}")
    return {
        "total": round(total_amount, 2),
        "por_gateway": gateways,
        "por_canal": channels
    }

def extract_resumen_liq(ws):
    data = {}
    for row in ws.iter_rows(min_row=1, max_row=50, values_only=True):
        if not row or not row[0]:
            continue
        label = str(row[0]).strip()
        nums = [to_num(v) for v in row[1:] if v is not None and isinstance(v,(int,float))]
        if label and nums:
            data[label] = nums[0]
    return data

def extract_reporte_2026(wb):
    result = {"meses": [], "por_tipo": []}

    for sn in wb.sheetnames:
        ws = wb[sn]
        if "hoja" in sn.lower() or sn.lower() == "sheet1":
            for row in ws.iter_rows(min_row=1, max_row=200, values_only=True):
                if not row or not row[0]:
                    continue
                label = str(row[0]).strip()
                nums = [to_num(v) for v in row[1:] if v is not None]
                if label and nums:
                    result["por_tipo"].append({"tipo": label, "valores": nums[:12]})
            break

    return result

def main():
    wb_liq = openpyxl.load_workbook(SOURCE_LIQ, data_only=True)
    wb_rep = openpyxl.load_workbook(SOURCE_REP, data_only=True)

    print(f"Liquidación hojas: {wb_liq.sheetnames[:8]}...")
    print(f"Reporte hojas: {wb_rep.sheetnames}")

    output = {
        "meta": {
            "fuentes": [SOURCE_LIQ.name, SOURCE_REP.name],
            "unidad": "M$ (miles de pesos chilenos)",
            "mes_transaccional": "Julio 2026"
        },
        "resumen_mes": {},
        "eventos": [],
        "pagos": {},
        "reporte_2026": {}
    }

    for sn in wb_liq.sheetnames:
        if "resumen" in sn.lower():
            print(f"  Resumen: {sn}")
            output["resumen_mes"] = extract_resumen_liq(wb_liq[sn])
            break

    for sn in wb_liq.sheetnames:
        if "ticket" in sn.lower() and "resumen" not in sn.lower():
            print(f"  Tickets: {sn}")
            output["eventos"] = extract_tickets(wb_liq[sn])
            break

    for sn in wb_liq.sheetnames:
        if "pago" in sn.lower():
            print(f"  Pagos: {sn}")
            output["pagos"] = extract_pagos(wb_liq[sn])
            break

    print(f"  Reporte 2026...")
    output["reporte_2026"] = extract_reporte_2026(wb_rep)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(output, f, ensure_ascii=False, indent=2)

    print(f"\nOK -> {OUT}")
    print(f"  Eventos: {len(output['eventos'])}")
    print(f"  Top 3 eventos: {[e['evento'][:40] for e in output['eventos'][:3]]}")

if __name__ == "__main__":
    main()
